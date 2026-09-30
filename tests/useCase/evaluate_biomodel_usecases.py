"""Evaluate the four selected BioModels top-1 runs against saved sources."""

from __future__ import annotations

import json
import argparse
import lzma
import pickle
import re
import sys
from collections import Counter, defaultdict
from pathlib import Path

import libsbml
import pandas as pd
from openpyxl import load_workbook

PROJECT_ROOT = Path(__file__).resolve().parents[2]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from core.database_search import load_chebi2kegg_dict, load_chebi_cleannames_dict
from core.reaction.hierarchy_relaxation import (
    load_chebi_parent_map,
    merge_chebi_to_kegg_mapping,
    normalize_chebi,
)


HERE = Path(__file__).resolve().parent
RESULTS = HERE / "results"
PUBLICATIONS = HERE / "context" / "publications"


def _norm(value: object) -> str:
    return re.sub(r"[^A-Z0-9]", "", str(value).upper())


def _name_without_charge_suffix(value: object) -> str:
    """Ignore only a trailing parenthesized formal charge in ChEBI labels."""
    return re.sub(r"\(\s*\d*\s*[+-]\s*\)\s*$", "", str(value)).strip()


def _annotation_id(value: object) -> str:
    return str(value).split(":", 1)[-1] if value else ""


def _save(model_id: str, metrics: dict, suffix: str = "") -> None:
    path = RESULTS / model_id / f"evaluation{suffix}.json"
    path.write_text(json.dumps(metrics, indent=2) + "\n", encoding="utf-8")
    print(path)


def _model_summary(model_id: str) -> dict[str, int]:
    document = libsbml.readSBMLFromFile(str(HERE / f"{model_id}.xml"))
    model = document.getModel()
    qual = model.getPlugin("qual")
    if qual and qual.getNumQualitativeSpecies():
        return {
            "species": qual.getNumQualitativeSpecies(),
            "reactions": qual.getNumTransitions(),
        }
    return {"species": model.getNumSpecies(), "reactions": model.getNumReactions()}


def _prediction_key(row: pd.Series) -> str:
    identity = row["identity"]
    if identity and ":" not in identity:
        return _norm(identity)
    return _norm(row["annotation_label"])


def _parse_table_s3() -> dict[str, str]:
    rows = {}
    path = HERE / "context" / "MODEL2501150001_TableS3.txt"
    for line in path.read_text(encoding="utf-8").splitlines():
        if ": " in line and not line.startswith("Species labels"):
            species_id, description = line.split(": ", 1)
            rows[species_id] = description
    return rows


def _table_s3_expected(species_id: str, description: str) -> set[str]:
    """Return the unique base identities asserted by the Table S3 description."""
    if _norm(species_id) == "SORAFENIB":
        return {"SORAFENIB"}
    text = f"{species_id} {description}"
    upper = text.upper()
    expected: set[str] = set()
    if "EGF-BOUND" in upper:
        expected.update({"EGF", "EGFR"})
    elif "EGFR" in upper:
        expected.add("EGFR")
    if "GRB2" in upper:
        expected.add("GRB2")
    if "SOS" in upper:
        expected.add("SOS")
    if "RAS" in upper:
        expected.add("RAS")
    if "GDP" in upper:
        expected.add("GDP")
    if "GTP" in upper:
        expected.add("GTP")
    if "BRAF" in upper:
        expected.add("BRAF")
    if "RAF1" in upper:
        expected.add("RAF1")
    if "MEK" in upper:
        expected.add("MEK")
    if "ERK" in upper:
        expected.add("ERK")
    if "SORAFENIB" in upper or re.search(r"I(?:B?RAF)", text, re.I):
        expected.add("SORAFENIB")
    return expected


def _model250115_identity(value: str) -> str:
    key = _norm(value)
    aliases = {
        "SOS1": "SOS",
        "SOS2": "SOS",
        "HRAS": "RAS",
        "KRAS": "RAS",
        "NRAS": "RAS",
        "MAP2K1": "MEK",
        "MAP2K2": "MEK",
        "MAPK1": "ERK",
        "MAPK3": "ERK",
    }
    return aliases.get(key, key)


def _evaluate_model250115_condition(filename: str) -> dict:
    df = pd.read_csv(RESULTS / "MODEL2501150001" / filename, dtype=str).fillna("")
    references = _parse_table_s3()
    groups = {_norm(key): group for key, group in df.groupby(df["id"].map(_norm), sort=False)}
    complete = 0
    recovered = 0
    total_expected = 0
    wrong = []
    missing = []
    unmatched_reference_species = []
    for species_id, description in references.items():
        if _norm(species_id) not in groups:
            unmatched_reference_species.append(species_id)
            continue
        expected = _table_s3_expected(species_id, description)
        group = groups[_norm(species_id)]
        observed = {
            _model250115_identity(_prediction_key(row))
            for _, row in group[group["annotation"] != ""].iterrows()
        }
        hits = expected & observed
        complete += expected <= observed
        recovered += len(hits)
        total_expected += len(expected)
        for identity in sorted(observed - expected):
            rows = group[
                group.apply(
                    lambda row: _model250115_identity(_prediction_key(row)) == identity,
                    axis=1,
                )
            ]
            wrong.append(
                {
                    "id": species_id,
                    "predicted_identity": identity,
                    "annotations": sorted(set(rows["annotation"]) - {""}),
                }
            )
        for identity in sorted(expected - observed):
            missing.append({"id": species_id, "missing_identity": identity})
    return {
        "species_with_candidates": int(df.loc[df["annotation"] != "", "id"].nunique()),
        "complex_species": int(df.loc[df["type"] == "complex", "id"].nunique()),
        "component_groups": int(
            df.loc[df["component_id"] != "", ["id", "component_id"]].drop_duplicates().shape[0]
        ),
        "reference_species_complete": complete,
        "reference_species_total": len(references) - len(unmatched_reference_species),
        "expected_identities_recovered": recovered,
        "expected_identities_total": total_expected,
        "wrong_predictions": wrong,
        "missing_identities": missing,
        "source_alignment": {
            "table_s3_species_absent_from_sbml": unmatched_reference_species,
            "sbml_species_absent_from_table_s3": sorted(
                key for key in df["id"].drop_duplicates()
                if _norm(key) not in {_norm(reference) for reference in references}
            ),
        },
    }


def evaluate_model2501150001(suffix: str = "") -> None:
    _save(
        "MODEL2501150001",
        {
            "model": _model_summary("MODEL2501150001"),
            "ranking": "top-1 identity per species or complex component",
            "sbml_only": _evaluate_model250115_condition(f"sbml_only{suffix}_species.csv"),
            "table_s3": _evaluate_model250115_condition(f"table_s3{suffix}_species.csv"),
        },
        suffix,
    )


def evaluate_model2503190002(suffix: str = "") -> None:
    model_id = "MODEL2503190002"
    df = pd.read_csv(RESULTS / model_id / f"full{suffix}_species.csv", dtype=str).fillna("")
    expected = {
        "ATM": "ATM", "ATR": "ATR", "p53": "TP53", "HDAC4": "HDAC4",
        "PP2A": "PPP2CA", "I2PP2A": "SET", "TOPBP1": "TOPBP1",
        "CK2": "CSNK2A1", "CHK1": "CHEK1", "CHK2": "CHEK2",
        "NRF2": "NFE2L2", "KEAP1": "KEAP1", "p21": "CDKN1A",
        "BAX": "BAX", "PUMA": "BBC3", "Omaveloxolone": "OMAVELOXOLONE",
        "Spermidine": "SPERMIDINE",
    }
    expected_by_id: dict[str, set[str]] = {}
    for species_id in df["id"].drop_duplicates():
        if species_id == "PP2A_I2PP2A":
            expected_by_id[species_id] = {"PPP2CA", "SET"}
            continue
        for token, identity in expected.items():
            if species_id == token or species_id.startswith(f"{token}_"):
                expected_by_id[species_id] = {_norm(identity)}
                break

    complete = 0
    wrong = []
    missing = []
    for species_id, reference in expected_by_id.items():
        group = df[df["id"] == species_id]
        observed = {_prediction_key(row) for _, row in group[group["annotation"] != ""].iterrows()}
        complete += observed == reference
        for identity in sorted(observed - reference):
            wrong.append({"id": species_id, "predicted_identity": identity})
        for identity in sorted(reference - observed):
            missing.append({"id": species_id, "missing_identity": identity})
    _save(
        model_id,
        {
            "model": _model_summary(model_id),
            "ranking": "top-1 identity per species or complex component",
            "species_with_candidates": int(df.loc[df["annotation"] != "", "id"].nunique()),
            "molecular_species_complete": complete,
            "molecular_species_total": len(expected_by_id),
            "wrong_predictions": wrong,
            "missing_identities": missing,
            "process_state_abstentions": ["DNA_damage", "Autophagy_active", "Autophagy_inactive"],
        },
        suffix,
    )


def _parse_dmms_metadata(path: Path) -> dict[str, dict[str, object]]:
    text = path.read_text(encoding="utf-8", errors="replace")
    metadata = {}
    for block in re.findall(r"NodeMetaData\{(.*?)NodeMetaData\}", text, re.S):
        def value(key: str) -> str:
            match = re.search(rf"^{key}:[ \t]*(.*)$", block, re.M)
            return match.group(1).strip() if match else ""

        name = value("NodeName")
        if name:
            metadata[name] = {
                "genes": [x.strip() for x in value("NodeGenes").split(",") if x.strip()],
                "type": value("NodeType"),
                "description": value("NodeDescription"),
                "notes": value("NodeNotes"),
            }
    return metadata


def evaluate_model2506050001(suffix: str = "") -> None:
    model_id = "MODEL2506050001"
    df = pd.read_csv(RESULTS / model_id / f"full{suffix}_species.csv", dtype=str).fillna("")
    metadata = _parse_dmms_metadata(PUBLICATIONS / model_id / "Hypoxia_EMT_Model.dmms")
    document = libsbml.readSBMLFromFile(str(HERE / f"{model_id}.xml"))
    qual = document.getModel().getPlugin("qual")
    id_to_name = {
        qual.getQualitativeSpecies(i).getId(): qual.getQualitativeSpecies(i).getName()
        for i in range(qual.getNumQualitativeSpecies())
    }
    taxon_by_gene: dict[str, str] = {}
    symbol_by_gene: dict[str, str] = {}
    ids_by_symbol: dict[str, set[str]] = defaultdict(set)
    for tax_id in ("9606", "10090", "10116"):
        path = PROJECT_ROOT / "data" / "ncbigene" / f"ncbigene2names_tax{tax_id}_protein-coding.lzma"
        with lzma.open(path, "rb") as handle:
            names = pickle.load(handle)
        for gene_id, gene_names in names.items():
            gene_id = str(gene_id)
            symbol = str(gene_names[0]) if gene_names else gene_id
            symbol_key = _norm(symbol)
            taxon_by_gene[gene_id] = tax_id
            symbol_by_gene[gene_id] = symbol
            ids_by_symbol[symbol_key].add(gene_id)

    reference_nodes = {
        species_id: {
            "name": name,
            "genes": set(metadata[name]["genes"]),
        }
        for species_id, name in id_to_name.items()
        if metadata[name]["genes"]
    }
    expected_occurrences = 0
    predicted_occurrences = 0
    exact_hits = 0
    nodes_any = 0
    nodes_complete = 0
    ortholog_supported = []
    related_nonreference = []
    wrong = []
    missing = []
    taxon_counts = {
        tax_id: {
            "source_expected": 0,
            "predicted": 0,
            "source_exact_hits": 0,
            "ortholog_supported": 0,
            "related_nonreference": 0,
            "wrong": 0,
        }
        for tax_id in ("9606", "10090", "10116")
    }
    related_identities = {"pAPC": {"ANAPC4", "ANAPC11"}}
    for species_id, reference in reference_nodes.items():
        expected_ids = reference["genes"]
        reference_genes = [
            {
                "entrez_id": gene_id,
                "name": symbol_by_gene.get(gene_id, "unknown"),
                "tax_id": taxon_by_gene.get(gene_id, "unknown"),
            }
            for gene_id in sorted(expected_ids, key=int)
        ]
        source_symbols = {
            _norm(symbol_by_gene[gene_id])
            for gene_id in expected_ids
            if gene_id in symbol_by_gene
        }
        accepted_ortholog_ids = {
            gene_id
            for symbol in source_symbols
            for gene_id in ids_by_symbol.get(symbol, set())
        }
        group = df[(df["id"] == species_id) & (df["annotation"] != "")]
        predicted_rows = [
            {
                "entrez_id": _annotation_id(row["annotation"]),
                "label": row["annotation_label"],
                "identity": row["identity"],
                "tax_id": row["tax_id"],
            }
            for _, row in group.iterrows()
        ]
        predicted_ids = {row["entrez_id"] for row in predicted_rows}
        hits = expected_ids & predicted_ids
        expected_occurrences += len(expected_ids)
        predicted_occurrences += len(predicted_ids)
        exact_hits += len(hits)
        nodes_any += bool(hits)
        nodes_complete += expected_ids <= predicted_ids
        for tax_id in taxon_counts:
            taxon_counts[tax_id]["source_expected"] += sum(
                taxon_by_gene.get(gene_id) == tax_id for gene_id in expected_ids
            )
            taxon_counts[tax_id]["predicted"] += sum(
                row["tax_id"] == tax_id for row in predicted_rows
            )
            taxon_counts[tax_id]["source_exact_hits"] += sum(
                gene_id in hits and taxon_by_gene.get(gene_id) == tax_id
                for gene_id in expected_ids
            )
        for row in predicted_rows:
            if row["entrez_id"] in expected_ids:
                continue
            assessment = {
                "id": species_id,
                "name": reference["name"],
                **row,
                "reference_genes": reference_genes,
            }
            tax_id = row["tax_id"]
            if row["entrez_id"] in accepted_ortholog_ids:
                ortholog_supported.append(assessment)
                taxon_counts[tax_id]["ortholog_supported"] += 1
            elif _norm(row["identity"]) in related_identities.get(reference["name"], set()):
                related_nonreference.append(assessment)
                taxon_counts[tax_id]["related_nonreference"] += 1
            else:
                wrong.append(assessment)
                taxon_counts[tax_id]["wrong"] += 1
        for gene_id in sorted(expected_ids - predicted_ids, key=int):
            missing.append(
                {
                    "id": species_id,
                    "name": reference["name"],
                    "entrez_id": gene_id,
                    "label": symbol_by_gene.get(gene_id, "unknown"),
                    "tax_id": taxon_by_gene.get(gene_id, "unknown"),
                }
            )

    abstract_predictions = []
    for species_id, name in id_to_name.items():
        if metadata[name]["genes"]:
            continue
        rows = df[(df["id"] == species_id) & (df["annotation"] != "")]
        if not rows.empty:
            abstract_predictions.append(
                {
                    "id": species_id,
                    "name": name,
                    "identities": sorted(set(rows["identity"])),
                    "entrez_ids": sorted({_annotation_id(x) for x in rows["annotation"]}, key=int),
                }
            )

    _save(
        model_id,
        {
            "model": _model_summary(model_id),
            "configuration": {
                "entity_type": "gene",
                "database": "NCBI Gene",
                "tax_ids": ["9606", "10090", "10116"],
                "ranking": "top-1 identity per species or complex component, then taxon expansion",
            },
            "species_with_candidates": int(df.loc[df["annotation"] != "", "id"].nunique()),
            "complex_species": int(df.loc[df["type"] == "complex", "id"].nunique()),
            "component_groups": int(
                df.loc[df["component_id"] != "", ["id", "component_id"]].drop_duplicates().shape[0]
            ),
            "reference_nodes": len(reference_nodes),
            "source_reference_entrez_occurrences": expected_occurrences,
            "predicted_entrez_occurrences_on_reference_nodes": predicted_occurrences,
            "source_exact_entrez_hits": exact_hits,
            "source_exact_entrez_recall": exact_hits / expected_occurrences,
            "source_exact_entrez_precision": exact_hits / predicted_occurrences,
            "verified_ortholog_predictions_not_listed_in_source": len(ortholog_supported),
            "reviewed_supported_predictions": exact_hits + len(ortholog_supported),
            "reviewed_supported_precision": (
                exact_hits + len(ortholog_supported)
            ) / predicted_occurrences,
            "reference_nodes_with_any_source_exact_hit": nodes_any,
            "reference_nodes_with_all_deposited_ids_recovered": nodes_complete,
            "per_taxon": taxon_counts,
            "ortholog_supported_predictions": ortholog_supported,
            "biologically_related_nonreference_predictions": related_nonreference,
            "wrong_predictions": wrong,
            "missing_references": missing,
            "nodes_without_source_gene_mapping": len(id_to_name) - len(reference_nodes),
            "predictions_on_nodes_without_source_gene_mapping": abstract_predictions,
        },
        suffix,
    )


def evaluate_model2507280001(suffix: str = "") -> None:
    model_id = "MODEL2507280001"
    df = pd.read_csv(RESULTS / model_id / f"full{suffix}_species.csv", dtype=str).fillna("")
    workbook = load_workbook(PUBLICATIONS / model_id / "supplement.xlsx", read_only=True, data_only=True)
    rows = list(workbook["Metabolites"].iter_rows(min_row=2, values_only=True))
    gold = {str(row[0]): str(row[5]) for row in rows if row[0] and row[5] not in (None, "", "NA")}
    reference_names = {str(row[0]): str(row[1]) for row in rows if row[0] and row[1]}
    chebi_to_kegg = merge_chebi_to_kegg_mapping(load_chebi2kegg_dict())
    chebi_names = load_chebi_cleannames_dict()
    parent_map = load_chebi_parent_map()
    local_kegg = {kegg for values in chebi_to_kegg.values() for kegg in values}
    scorable = {species_id: kegg for species_id, kegg in gold.items() if kegg in local_kegg}
    predictions = {
        species_id: next((value for value in group["annotation"] if value.startswith("CHEBI:")), "")
        for species_id, group in df.groupby("id", sort=False)
    }
    prediction_labels = {
        species_id: next(
            (row["annotation_label"] for _, row in group.iterrows()
             if str(row["annotation"]).startswith("CHEBI:")),
            "",
        )
        for species_id, group in df.groupby("id", sort=False)
    }

    exact_hits = []
    relaxed_hits = []
    for species_id, kegg in scorable.items():
        annotation = predictions.get(species_id, "")
        exact = annotation and kegg in normalize_chebi(
            annotation, chebi_to_kegg, parent_map, level=0, max_depth=0
        )
        relaxed = annotation and kegg in normalize_chebi(
            annotation, chebi_to_kegg, parent_map, level=2, max_depth=2
        )
        if exact:
            exact_hits.append(species_id)
        if relaxed:
            relaxed_hits.append(species_id)

    crossref_identity_review_ids = {
        "M_4pasp_c", "M_thdp_c", "M_succ_c", "M_succ_p", "M_succ_s", "M_succ_e",
        "M_fe2_c", "M_fe2_p", "M_fe2_s", "M_fe2_e", "M_cu2_c", "M_cu2_p",
        "M_cu2_e", "M_mn2_c", "M_mn2_p", "M_mn2_s", "M_mn2_e", "M_dhor__S_c",
    }
    additional_identity_review_ids = {
        "M_g6p_c", "M_nad_c", "M_13dpg_c", "M_actp_c", "M_6pgc_c",
        "M_asp__L_c", "M_asp__L_p", "M_asp__L_s", "M_asp__L_e",
        "M_mlthf_c", "M_mlthf_p", "M_mlthf_s", "M_mlthf_e",
        "M_so3_c", "M_chor_c", "M_3c3hmp_c",
        "M_5mtr_c", "M_5mtr_p", "M_5mtr_s", "M_5mtr_e",
        "M_glu__D_c", "M_ACP_c", "M_orot_c", "M_ipdp_c",
        "M_pnto__R_c", "M_pnto__R_p", "M_pnto__R_s", "M_pnto__R_e",
        "M_amob_c", "M_uaccg_c", "M_uamr_c", "M_uama_c", "M_cu2_s", "M_thmpp_s",
    }
    reviewed_ids = crossref_identity_review_ids | additional_identity_review_ids
    # Manual identity review belongs to a species–prediction pair, not the
    # species ID alone. A new validation run must not inherit approval for a
    # different ChEBI term merely because it annotates the same SBML species.
    approved_annotations = {}
    if suffix:
        baseline_review_path = RESULTS / model_id / "evaluation.json"
        if baseline_review_path.exists():
            baseline_review = json.loads(baseline_review_path.read_text())
            approved_annotations = {
                row["id"]: row["predicted_annotation"]
                for row in baseline_review["reviewed_identity_matches_outside_automated_mapping"]
            }
    reviewed_hits = sorted(
        species_id for species_id in reviewed_ids
        if species_id in scorable and species_id not in relaxed_hits
        and predictions.get(species_id)
        and (not suffix or predictions[species_id] == approved_annotations.get(species_id))
    )
    charge_label_hits = sorted(
        species_id for species_id in scorable
        if species_id not in relaxed_hits and species_id not in reviewed_hits
        and prediction_labels.get(species_id)
        and _norm(_name_without_charge_suffix(prediction_labels[species_id]))
        == _norm(reference_names.get(species_id, ""))
    )
    chebi_synonym_hits = sorted(
        species_id for species_id in scorable
        if species_id not in relaxed_hits and species_id not in reviewed_hits
        and species_id not in charge_label_hits and predictions.get(species_id)
        and _annotation_id(predictions[species_id]) in chebi_names.get(
            re.sub(r"[^a-z0-9]", "", reference_names.get(species_id, "").lower()), []
        )
    )
    # Publication-name/ChEBI-ontology review of the remaining validation-only
    # pairs. Lock both species and accession so later runs are not implicitly
    # approved if their prediction changes.
    manually_supported_validation_pairs = {
        ("M_nadh_c", "CHEBI:57945"),
        ("M_nadph_c", "CHEBI:57783"),
        ("M_h2mb4p_c", "CHEBI:128753"),
        ("M_pheme_c", "CHEBI:60344"),
        ("M_pheme_p", "CHEBI:60344"),
        ("M_pheme_s", "CHEBI:60344"),
        ("M_pheme_e", "CHEBI:60344"),
        ("M_pg160_e", "CHEBI:85270"),
        # ChEBI:57791 is a tautomer of meso-DAP (ChEBI:16488);
        # ChEBI:57328 is the conjugate base of dephospho-CoA (ChEBI:15468).
        ("M_26dap__M_c", "CHEBI:57791"),
        ("M_dpcoa_c", "CHEBI:57328"),
    }
    manual_validation_hits = sorted(
        species_id for species_id in scorable
        if suffix and (species_id, predictions.get(species_id, ""))
        in manually_supported_validation_pairs
        and species_id not in relaxed_hits and species_id not in reviewed_hits
        and species_id not in charge_label_hits and species_id not in chebi_synonym_hits
    )
    known_incomplete_validation_pairs = {
        ("M_3mop_c", "CHEBI:28654"),
        ("M_3mop_s", "CHEBI:28654"),
        ("M_3hhexACP_c", "CHEBI:37035"),
        ("M_3ooctACP_c", "CHEBI:44680"),
        ("M_3hddecACP_c", "CHEBI:76616"),
        ("M_myrsACP_c", "CHEBI:28875"),
        ("M_nadph_c", "CHEBI:13392"),  # NAD(P)H also includes NADH, not just NADPH.
    }
    reviewed_predictions = []
    wrong = []
    unresolved = []
    no_candidate = []
    incomplete_ids = {
        "M_nadph_c", "M_actACP_c", "M_3oddecACP_c", "M_3hoctaACP_c", "M_mlhpglu_c",
    }
    for species_id, kegg in scorable.items():
        annotation = predictions.get(species_id, "")
        if not annotation:
            no_candidate.append(
                {
                    "id": species_id,
                    "reference_name": reference_names.get(species_id, ""),
                    "reference_kegg": kegg,
                }
            )
        elif (species_id in reviewed_hits or species_id in charge_label_hits
              or species_id in chebi_synonym_hits or species_id in manual_validation_hits):
            row = df[df["id"] == species_id].iloc[0]
            reviewed_predictions.append(
                {
                    "id": species_id,
                    "predicted_annotation": annotation,
                    "predicted_label": row["annotation_label"],
                    "reference_name": reference_names.get(species_id, ""),
                    "reference_kegg": kegg,
                    "basis": (
                        "charge_normalized_exact_label"
                        if species_id in charge_label_hits else (
                            "chebi_reference_name_synonym"
                            if species_id in chebi_synonym_hits else "manual_identity_review"
                        )
                    ),
                }
            )
        elif species_id not in relaxed_hits:
            row = df[df["id"] == species_id].iloc[0]
            candidate_record = (
                {
                    "id": species_id,
                    "predicted_annotation": annotation,
                    "predicted_label": row["annotation_label"],
                    "reference_name": reference_names.get(species_id, ""),
                    "reference_kegg": kegg,
                }
            )
            if suffix and (species_id, annotation) in known_incomplete_validation_pairs:
                candidate_record["assessment"] = "incomplete_identity"
                wrong.append(candidate_record)
            elif suffix:
                unresolved.append(candidate_record)
            else:
                candidate_record["assessment"] = (
                    "incomplete_or_overbroad" if species_id in incomplete_ids else "wrong_identity"
                )
                wrong.append(candidate_record)
    document = libsbml.readSBMLFromFile(str(HERE / f"{model_id}.xml"))
    sbml_ids = {
        document.getModel().getSpecies(i).getId()
        for i in range(document.getModel().getNumSpecies())
    }
    supplement_ids = {str(row[0]) for row in rows if row[0]}
    supported_total = len(
        set(relaxed_hits) | set(reviewed_hits) | set(charge_label_hits)
        | set(chebi_synonym_hits) | set(manual_validation_hits)
    )
    _save(
        model_id,
        {
            "model": _model_summary(model_id),
            "ranking": "top-1 identity per species",
            "species_with_candidates": int(df.loc[df["annotation"] != "", "id"].nunique()),
            "supplement_kegg_references": len(gold),
            "locally_scorable_kegg_references": len(scorable),
            "exact_top1_hits": len(exact_hits),
            "exact_top1_rate": len(exact_hits) / len(scorable),
            "two_hop_top1_hits": len(relaxed_hits),
            "two_hop_top1_rate": len(relaxed_hits) / len(scorable),
            "charge_normalized_exact_name_hits": len(charge_label_hits),
            "chebi_reference_name_synonym_hits": len(chebi_synonym_hits),
            "manually_supported_validation_pairs": len(manual_validation_hits),
            "reviewed_identity_matches_outside_automated_mapping": reviewed_predictions,
            "reviewed_supported_top1": supported_total,
            "reviewed_supported_top1_rate": supported_total / len(scorable),
            "wrong_predictions": wrong,
            "unresolved_predictions": unresolved,
            "no_candidate_reference_species": no_candidate,
            "unscorable_reference_kegg_ids": dict(Counter(gold.values()) - Counter(scorable.values())),
            "supplement_rows": len(rows),
            "literal_sbml_id_matches": len(sbml_ids & supplement_ids),
            "source_alignment_notes": {
                "normalized_biomass_pool_ids": 3,
                "sbml_species_without_supplement_row": ["M_h_m_p"],
            },
        },
        suffix,
    )


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--validation",
        action="store_true",
        help="Evaluate *_validation_species.csv and save evaluation_validation.json.",
    )
    parser.add_argument("--source-fidelity", action="store_true",
                        help="Evaluate *_source_fidelity_species.csv separately.")
    args = parser.parse_args()
    suffix = ("_source_fidelity" if args.source_fidelity else
              "_validation" if args.validation else "")
    evaluate_model2501150001(suffix)
    evaluate_model2503190002(suffix)
    evaluate_model2506050001(suffix)
    evaluate_model2507280001(suffix)


if __name__ == "__main__":
    main()
